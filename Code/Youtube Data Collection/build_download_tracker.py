#!/usr/bin/env python3
"""Build Excel and CSV trackers from completed channel downloads."""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

MEDIA_EXTENSIONS = {
    ".avi", ".flac", ".m4a", ".mkv", ".mov", ".mp3",
    ".mp4", ".ogg", ".opus", ".wav", ".webm",
}
VIDEO_ID_PATTERN = re.compile(r"\[([A-Za-z0-9_-]{11})\]")
FORMAT_STREAM_PATTERN = re.compile(r"\.f\d+\.[^.]+$", re.IGNORECASE)
CHANNEL_ID_PATTERN = re.compile(r"/channel/([^/?#]+)", re.IGNORECASE)
CHALLENGE_DAY_PATTERN = re.compile(r"\bday[\s:#-]*(\d{1,3})\b", re.IGNORECASE)


def parse_args() -> argparse.Namespace:
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(
        description="Create Excel and CSV trackers for completed channel downloads."
    )
    parser.add_argument(
        "-o", "--output", type=Path, default=Path("channel_downloads"),
        help="Channel download directory (default: ./channel_downloads)",
    )
    parser.add_argument(
        "--channels", type=Path, default=script_dir / "channels_75_hard.tsv",
        help="Original TSV channel list used for the vlogger summary",
    )
    return parser.parse_args()


def normalize_upload_date(value: object) -> str:
    text = str(value or "").strip()
    if re.fullmatch(r"\d{8}", text):
        return f"{text[:4]}-{text[4:6]}-{text[6:8]}"
    return text


def parse_iso_date(value: object) -> date | None:
    try:
        return date.fromisoformat(str(value))
    except (TypeError, ValueError):
        return None


def extract_challenge_day(title: str) -> int | None:
    match = CHALLENGE_DAY_PATTERN.search(title)
    if not match:
        return None
    challenge_day = int(match.group(1))
    return challenge_day if 1 <= challenge_day <= 75 else None


def completed_media_by_id(output_dir: Path) -> dict[str, Path]:
    media: dict[str, Path] = {}
    if not output_dir.exists():
        return media

    for folder_name in ("videos", "audio"):
        for path in output_dir.rglob(f"{folder_name}/*"):
            if not path.is_file() or path.suffix.lower() not in MEDIA_EXTENSIONS:
                continue
            if path.name.endswith(".part") or FORMAT_STREAM_PATTERN.search(path.name):
                continue
            match = VIDEO_ID_PATTERN.search(path.name)
            if match:
                media.setdefault(match.group(1), path)
    return media


def read_download_records(output_dir: Path) -> list[dict[str, object]]:
    media_by_id = completed_media_by_id(output_dir)
    records_by_id: dict[str, dict[str, object]] = {}
    metadata_files = (
        sorted(output_dir.rglob("metadata/*.info.json"))
        if output_dir.exists() else []
    )

    for metadata_path in metadata_files:
        try:
            data = json.loads(metadata_path.read_text(encoding="utf-8"))
        except (OSError, UnicodeDecodeError, json.JSONDecodeError) as error:
            print(f"Warning: could not read {metadata_path}: {error}")
            continue

        video_id = str(data.get("id") or "").strip()
        media_path = media_by_id.get(video_id)
        if not video_id or media_path is None:
            continue

        try:
            media_stat = media_path.stat()
            downloaded_at = datetime.fromtimestamp(
                media_stat.st_mtime
            ).astimezone().isoformat(timespec="seconds")
            file_size_bytes = media_stat.st_size
        except OSError:
            downloaded_at = ""
            file_size_bytes = 0

        title = str(data.get("title") or "")
        duration_seconds = float(data.get("duration") or 0)
        records_by_id[video_id] = {
            "Vlogger": str(data.get("channel") or data.get("uploader") or "Unknown Channel"),
            "Video Title": title,
            "Challenge Day": extract_challenge_day(title) or "",
            "Video ID": video_id,
            "Upload Date": normalize_upload_date(data.get("upload_date")),
            "Duration Minutes": round(duration_seconds / 60, 2),
            "File Size MiB": round(file_size_bytes / (1024 * 1024), 2),
            "Video URL": str(
                data.get("webpage_url")
                or f"https://www.youtube.com/watch?v={video_id}"
            ),
            "Channel ID": str(data.get("channel_id") or ""),
            "Downloaded At": downloaded_at,
            "Media File": str(media_path.relative_to(output_dir)),
        }

    return sorted(
        records_by_id.values(),
        key=lambda row: (
            str(row["Vlogger"]).casefold(),
            str(row["Upload Date"]),
            str(row["Video Title"]).casefold(),
        ),
    )


def read_listed_channels(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        return []

    channels: list[dict[str, str]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file, delimiter="\t")
        for row in reader:
            name = (row.get("Channel Name") or "").strip()
            url = (row.get("Channel URL") or "").strip()
            if not name or not url:
                continue
            match = CHANNEL_ID_PATTERN.search(url)
            channels.append({
                "Vlogger": name,
                "Channel URL": url,
                "Channel ID": match.group(1) if match else "",
            })
    return channels


def summarize_matches(
    vlogger: str,
    channel_url: str,
    matches: list[dict[str, object]],
) -> dict[str, object]:
    upload_dates = sorted({
        parsed for record in matches
        if (parsed := parse_iso_date(record["Upload Date"])) is not None
    })
    challenge_days = sorted({
        int(record["Challenge Day"])
        for record in matches if record["Challenge Day"] != ""
    })
    missing_days = [day for day in range(1, 76) if day not in challenge_days]
    total_duration_minutes = sum(float(record["Duration Minutes"]) for record in matches)
    total_file_mib = sum(float(record["File Size MiB"]) for record in matches)
    unique_upload_days = len(upload_dates)

    return {
        "Vlogger": vlogger,
        "Channel URL": channel_url,
        "Videos Downloaded": len(matches),
        "Unique Upload Days": unique_upload_days,
        "First Upload Date": upload_dates[0].isoformat() if upload_dates else "",
        "Last Upload Date": upload_dates[-1].isoformat() if upload_dates else "",
        "Calendar Span Days": (
            (upload_dates[-1] - upload_dates[0]).days + 1 if upload_dates else 0
        ),
        "Videos per Upload Day": (
            round(len(matches) / unique_upload_days, 2) if unique_upload_days else 0
        ),
        "Challenge Days Found": len(challenge_days),
        "First Challenge Day": challenge_days[0] if challenge_days else "",
        "Last Challenge Day": challenge_days[-1] if challenge_days else "",
        "75-Day Coverage": len(challenge_days) / 75,
        "Missing Challenge Days": ", ".join(map(str, missing_days)),
        "Total Duration Hours": round(total_duration_minutes / 60, 2),
        "Average Video Minutes": (
            round(total_duration_minutes / len(matches), 2) if matches else 0
        ),
        "Total File Size GB": round(total_file_mib / 1024, 3),
        "Downloaded Video Titles": " | ".join(sorted(
            {str(record["Video Title"]) for record in matches}, key=str.casefold
        )),
    }


def build_summary_rows(
    records: list[dict[str, object]],
    listed_channels: list[dict[str, str]],
) -> list[dict[str, object]]:
    records_by_channel_id: dict[str, list[dict[str, object]]] = defaultdict(list)
    records_by_name: dict[str, list[dict[str, object]]] = defaultdict(list)
    for record in records:
        channel_id = str(record["Channel ID"])
        if channel_id:
            records_by_channel_id[channel_id].append(record)
        records_by_name[str(record["Vlogger"]).casefold()].append(record)

    summary_rows: list[dict[str, object]] = []
    if listed_channels:
        for channel in listed_channels:
            matches = records_by_channel_id.get(channel["Channel ID"], [])
            if not matches:
                matches = records_by_name.get(channel["Vlogger"].casefold(), [])
            summary_rows.append(summarize_matches(
                channel["Vlogger"], channel["Channel URL"], matches
            ))
    else:
        for _, matches in sorted(records_by_name.items()):
            summary_rows.append(summarize_matches(
                str(matches[0]["Vlogger"]), "", matches
            ))
    return summary_rows


def write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_name(f".{path.stem}.tmp{path.suffix}")
    with temporary_path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    temporary_path.replace(path)


def write_excel(
    path: Path,
    summary_rows: list[dict[str, object]],
    records: list[dict[str, object]],
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as error:
        raise RuntimeError(
            "Excel support requires openpyxl. Install requirements with: "
            "python3 -m pip install --upgrade -r requirements.txt"
        ) from error

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Vlogger Summary"
    detail_sheet = workbook.create_sheet("Downloaded Videos")
    dark_blue = "1F4E78"
    white = "FFFFFF"
    font_name = "Arial"
    generated_at = datetime.now().astimezone().replace(microsecond=0)

    summary_fields = [
        "Vlogger", "Channel URL", "Videos Downloaded", "Unique Upload Days",
        "First Upload Date", "Last Upload Date", "Calendar Span Days",
        "Videos per Upload Day", "Challenge Days Found", "First Challenge Day",
        "Last Challenge Day", "75-Day Coverage", "Missing Challenge Days",
        "Total Duration Hours", "Average Video Minutes", "Total File Size GB",
    ]
    summary_sheet["A1"] = "75 Hard Vlogger Tracker"
    summary_sheet["A1"].font = Font(name=font_name, size=16, bold=True, color=dark_blue)
    summary_sheet["A2"] = "Updated"
    summary_sheet["B2"] = generated_at.replace(tzinfo=None)
    summary_sheet["B2"].number_format = "mm/dd/yy h:mm AM/PM"
    summary_sheet["D2"] = "Vloggers listed"
    summary_sheet["E2"] = len(summary_rows)
    summary_sheet["G2"] = "With downloads"
    summary_sheet["H2"] = sum(1 for row in summary_rows if row["Videos Downloaded"])
    summary_sheet["J2"] = "Videos downloaded"
    summary_sheet["K2"] = len(records)
    summary_sheet["A3"] = (
        "Challenge coverage uses distinct DAY 1–75 numbers found in video titles. "
        "Upload-day metrics use distinct YouTube upload dates."
    )
    summary_sheet.merge_cells("A3:P3")
    summary_sheet["A3"].font = Font(name=font_name, size=10, italic=True, color="595959")

    for column, field in enumerate(summary_fields, start=1):
        cell = summary_sheet.cell(row=5, column=column, value=field)
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(name=font_name, size=10, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_number, row in enumerate(summary_rows, start=6):
        for column, field in enumerate(summary_fields, start=1):
            value = row[field]
            if field in {"First Upload Date", "Last Upload Date"} and value:
                value = parse_iso_date(value)
            cell = summary_sheet.cell(row=row_number, column=column, value=value)
            cell.font = Font(name=font_name, size=10)
            cell.alignment = Alignment(vertical="center")
            if field == "Channel URL" and value:
                cell.hyperlink = str(value)
                cell.style = "Hyperlink"
            elif field in {"First Upload Date", "Last Upload Date"}:
                cell.number_format = "mm/dd/yy"
            elif field == "75-Day Coverage":
                cell.number_format = "0.0%"
            elif field in {
                "Videos per Upload Day", "Total Duration Hours", "Average Video Minutes"
            }:
                cell.number_format = "0.00"
            elif field == "Total File Size GB":
                cell.number_format = "0.000"

    summary_sheet.freeze_panes = "A6"
    summary_sheet.sheet_view.showGridLines = False
    summary_sheet.auto_filter.ref = f"A5:P{max(5, 5 + len(summary_rows))}"
    summary_sheet.row_dimensions[1].height = 24
    summary_sheet.row_dimensions[5].height = 34
    summary_widths = {
        "A": 25, "B": 46, "C": 17, "D": 18, "E": 15, "F": 15,
        "G": 18, "H": 19, "I": 19, "J": 18, "K": 18, "L": 16,
        "M": 42, "N": 20, "O": 22, "P": 18,
    }
    for column, width in summary_widths.items():
        summary_sheet.column_dimensions[column].width = width
    if summary_rows:
        summary_table = Table(
            displayName="VloggerSummaryTable", ref=f"A5:P{5 + len(summary_rows)}"
        )
        summary_table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
        )
        summary_sheet.add_table(summary_table)
        summary_sheet.conditional_formatting.add(
            f"L6:L{5 + len(summary_rows)}",
            ColorScaleRule(
                start_type="num", start_value=0, start_color="FEE2E2",
                mid_type="num", mid_value=0.5, mid_color="FEF3C7",
                end_type="num", end_value=1, end_color="DCFCE7",
            ),
        )

    detail_fields = [
        "Vlogger", "Video Title", "Challenge Day", "Upload Date",
        "Duration Minutes", "File Size MiB", "Video ID", "Video URL",
        "Channel ID", "Downloaded At", "Media File",
    ]
    detail_sheet["A1"] = "Downloaded Videos"
    detail_sheet["A1"].font = Font(name=font_name, size=16, bold=True, color=dark_blue)
    detail_sheet["A2"] = (
        "Only completed media files are included. Partial format streams and failed downloads are excluded."
    )
    detail_sheet.merge_cells("A2:K2")
    detail_sheet["A2"].font = Font(name=font_name, size=10, italic=True, color="595959")
    for column, field in enumerate(detail_fields, start=1):
        cell = detail_sheet.cell(row=4, column=column, value=field)
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(name=font_name, size=10, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_number, record in enumerate(records, start=5):
        for column, field in enumerate(detail_fields, start=1):
            value = record[field]
            if field == "Upload Date" and value:
                value = parse_iso_date(value)
            elif field == "Downloaded At" and value:
                try:
                    value = datetime.fromisoformat(str(value)).replace(tzinfo=None)
                except ValueError:
                    pass
            cell = detail_sheet.cell(row=row_number, column=column, value=value)
            cell.font = Font(name=font_name, size=10)
            cell.alignment = Alignment(vertical="center")
            if field == "Video URL" and value:
                cell.hyperlink = str(value)
                cell.style = "Hyperlink"
            elif field == "Upload Date":
                cell.number_format = "mm/dd/yy"
            elif field == "Downloaded At":
                cell.number_format = "mm/dd/yy h:mm AM/PM"
            elif field in {"Duration Minutes", "File Size MiB"}:
                cell.number_format = "0.00"

    detail_sheet.freeze_panes = "A5"
    detail_sheet.sheet_view.showGridLines = False
    detail_sheet.auto_filter.ref = f"A4:K{max(4, 4 + len(records))}"
    detail_sheet.row_dimensions[1].height = 24
    detail_sheet.row_dimensions[4].height = 30
    detail_widths = {
        "A": 25, "B": 48, "C": 15, "D": 14, "E": 18, "F": 15,
        "G": 15, "H": 46, "I": 28, "J": 22, "K": 62,
    }
    for column, width in detail_widths.items():
        detail_sheet.column_dimensions[column].width = width
    if records:
        detail_table = Table(
            displayName="DownloadedVideosTable", ref=f"A4:K{4 + len(records)}"
        )
        detail_table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
        )
        detail_sheet.add_table(detail_table)

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_name(f".{path.name}.tmp")
    workbook.save(temporary_path)
    temporary_path.replace(path)


def build_trackers(output_dir: Path, channels_path: Path) -> tuple[int, int]:
    output_dir = output_dir.expanduser().resolve()
    channels_path = channels_path.expanduser().resolve()
    records = read_download_records(output_dir)
    listed_channels = read_listed_channels(channels_path)
    summary_rows = build_summary_rows(records, listed_channels)

    tracker_fields = [
        "Vlogger", "Video Title", "Challenge Day", "Video ID", "Upload Date",
        "Duration Minutes", "File Size MiB", "Video URL", "Channel ID",
        "Downloaded At", "Media File",
    ]
    write_csv(output_dir / "download_tracker.csv", tracker_fields, records)

    summary_fields = [
        "Vlogger", "Channel URL", "Videos Downloaded", "Unique Upload Days",
        "First Upload Date", "Last Upload Date", "Calendar Span Days",
        "Videos per Upload Day", "Challenge Days Found", "First Challenge Day",
        "Last Challenge Day", "75-Day Coverage", "Missing Challenge Days",
        "Total Duration Hours", "Average Video Minutes", "Total File Size GB",
        "Downloaded Video Titles",
    ]
    write_csv(output_dir / "vlogger_summary.csv", summary_fields, summary_rows)
    write_excel(output_dir / "vlogger_tracker.xlsx", summary_rows, records)
    return len(records), len(summary_rows)


def main() -> int:
    args = parse_args()
    try:
        video_count, vlogger_count = build_trackers(args.output, args.channels)
    except RuntimeError as error:
        print(f"Tracker error: {error}")
        return 1
    output_dir = args.output.expanduser().resolve()
    print(f"Excel tracker updated: {output_dir / 'vlogger_tracker.xlsx'}")
    print(f"Video CSV updated: {output_dir / 'download_tracker.csv'}")
    print(f"Vlogger CSV updated: {output_dir / 'vlogger_summary.csv'}")
    print(f"Completed videos tracked: {video_count}")
    print(f"Vloggers summarized: {vlogger_count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
